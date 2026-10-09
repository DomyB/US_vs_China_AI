"""Regenerate the readable workbook of the statements dataset from the CSV.

Usage: python scripts/statements_workbook.py [csv] [workbook]

The workbook (default data/manual/statements/political_statements_minerals.xlsx) is the owner's formatted copy of the
dataset: a README sheet, a readable `Statements` sheet (country names, labels without underscores, the source as a link,
sorted by date), a `Coverage` sheet of COUNTIFS/SUMPRODUCT formulas and the raw `Data` sheet. This script keeps the
workbook's layout and styles and rewrites the rows from the CSV, so the file stays in step with the dataset after
records are added: Data and Statements are rewritten in full (styles copied from the first data row), Coverage is
rebuilt section by section from the template's own headers and formulas with one row per country and mineral present,
and every formula range is widened to the new last row. No value is computed here that is not in the CSV.
"""
from __future__ import annotations

import re
import sys
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
CSV = ROOT / "data" / "manual" / "statements" / "political_statements_minerals.csv"
XLSX = ROOT / "data" / "manual" / "statements" / "political_statements_minerals.xlsx"
IN_SCOPE = {"ARG", "BOL", "BRA", "CHL", "COL", "ECU", "GUY", "PER", "PRY", "SUR", "URY", "VEN"}
COUNTRY = {"ARG": "Argentina", "BOL": "Bolivia", "BRA": "Brazil", "CHL": "Chile", "COL": "Colombia", "ECU": "Ecuador", "GUY": "Guyana", "PER": "Peru",
           "PRY": "Paraguay", "SUR": "Suriname", "URY": "Uruguay", "VEN": "Venezuela", "MEX": "Mexico", "PAN": "Panama", "REG": "Regional"}
READABLE = ["ID", "Date", "Country", "In 12-country scope", "Speaker", "Role", "Speaker bloc", "Channel", "Event / context", "Summary", "Minerals", "Themes",
            "Stance on China", "Stance on US", "Quote (original)", "Quote (English)", "Language", "Source", "Source type", "Verification", "Notes"]


def label(v: str) -> str:
    """`labor_social_conflict;foreign_investment` -> `labor social conflict, foreign investment`."""
    return ", ".join(part.strip().replace("_", " ") for part in str(v).split(";") if part.strip())


def copy_style(src, dst) -> None:
    dst.font, dst.fill, dst.border, dst.alignment, dst.number_format, dst.protection = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment), src.number_format, copy(src.protection)


def rewrite_rows(ws, rows: list[list], hyperlinks: dict[int, list[str | None]] | None = None) -> int:
    """Replace every data row of a sheet (row 2 onward) with `rows`, copying the first data row's cell styles."""
    styles = [copy_style_src(ws.cell(row=2, column=c)) for c in range(1, ws.max_column + 1)]
    height = ws.row_dimensions[2].height
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    for i, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=value)
            copy_style(styles[c - 1], cell)
            if hyperlinks and c in hyperlinks and hyperlinks[c][i - 2]:
                cell.hyperlink = hyperlinks[c][i - 2]
        if height:
            ws.row_dimensions[i].height = height
    last = len(rows) + 1
    for t in ws.tables.values():
        t.ref = f"A1:{get_column_letter(ws.max_column)}{last}"
    return last


def copy_style_src(cell):
    """A detached copy of a cell's style (the cell itself is deleted with its row)."""
    holder = openpyxl.cell.cell.Cell(cell.parent, row=cell.row, column=cell.column)
    copy_style(cell, holder)
    return holder


def widen(formula: str, old_last: int, new_last: int) -> str:
    return re.sub(rf"\${old_last}\b", f"${new_last}", formula)


def rebuild_coverage(ws, df: pd.DataFrame, old_last: int, new_last: int) -> None:
    """Rebuild the Coverage sheet from its own sections: each has a title row, a header row, data rows whose formulas
    reference their own row in column A, a Total row and (optionally) a Check row. Country sections get one row per
    country present (by count), the mineral section one row per tag present (template order first)."""
    max_row = ws.max_row
    sections = []
    r = 1
    while r <= max_row:
        a = ws.cell(row=r, column=1).value
        if a and ws.cell(row=r + 1, column=1).value in ("Code", "Tag"):
            title, header = r, r + 1
            rows = []
            k = header + 1
            while ws.cell(row=k, column=1).value not in (None, "") and ws.cell(row=k, column=2).value != "Total":
                rows.append(k)
                k += 1
            total = k if ws.cell(row=k, column=2).value == "Total" else None
            check = k + 1 if total and ws.cell(row=k + 1, column=2).value and str(ws.cell(row=k + 1, column=2).value).startswith("Check") else None
            sections.append({"title": title, "header": header, "rows": rows, "total": total, "check": check, "kind": ws.cell(row=header, column=1).value})
            r = (check or total or rows[-1]) + 1
        else:
            r += 1
    assert sections, "no sections recognised in the Coverage sheet"
    top = [(rr, [copy_style_src(ws.cell(row=rr, column=c)) for c in range(1, ws.max_column + 1)], [ws.cell(row=rr, column=c).value for c in range(1, ws.max_column + 1)]) for rr in range(1, sections[0]["title"])]
    snap = []
    for s in sections:
        snap.append({
            "kind": s["kind"],
            "title": (ws.cell(row=s["title"], column=1).value, copy_style_src(ws.cell(row=s["title"], column=1))),
            "header": [(ws.cell(row=s["header"], column=c).value, copy_style_src(ws.cell(row=s["header"], column=c))) for c in range(1, ws.max_column + 1)],
            "row": [(ws.cell(row=s["rows"][0], column=c).value, copy_style_src(ws.cell(row=s["rows"][0], column=c))) for c in range(1, ws.max_column + 1)],
            "codes": [ws.cell(row=rr, column=1).value for rr in s["rows"]],
            "total": [(ws.cell(row=s["total"], column=c).value, copy_style_src(ws.cell(row=s["total"], column=c))) for c in range(1, ws.max_column + 1)] if s["total"] else None,
            "check": [(ws.cell(row=s["check"], column=c).value, copy_style_src(ws.cell(row=s["check"], column=c))) for c in range(1, ws.max_column + 1)] if s["check"] else None,
            "first_row": s["rows"][0],
        })
    ws.delete_rows(1, ws.max_row)
    counts = df["country"].value_counts()
    countries = [c for c in counts.index]  # by count, ties in first-seen order
    minerals = []
    for v in df["minerals"]:
        for tag in str(v).split(";"):
            if tag.strip() and tag.strip() not in minerals:
                minerals.append(tag.strip())
    r = 1
    for rr, styles, values in top:
        for c, (st, v) in enumerate(zip(styles, values, strict=True), start=1):
            cell = ws.cell(row=rr, column=c, value=v)
            copy_style(st, cell)
        r = rr + 1
    r += 1
    for s in snap:
        cell = ws.cell(row=r, column=1, value=s["title"][0])
        copy_style(s["title"][1], cell)
        r += 1
        for c, (v, st) in enumerate(s["header"], start=1):
            cell = ws.cell(row=r, column=c, value=v)
            copy_style(st, cell)
        r += 1
        is_mineral = s["kind"] == "Tag"
        if is_mineral:
            codes = [m for m in s["codes"] if m in minerals] + [m for m in minerals if m not in s["codes"]]
        else:
            codes = countries
        first = r
        for code in codes:
            for c, (v, st) in enumerate(s["row"], start=1):
                if c == 1:
                    val = code
                elif c == 2:
                    val = code.replace("_", " ") if is_mineral else COUNTRY.get(code, code)
                elif isinstance(v, str) and v.startswith("="):
                    val = widen(re.sub(rf"\$A{s['first_row']}\b", f"$A{r}", v), old_last, new_last)
                else:
                    val = v
                cell = ws.cell(row=r, column=c, value=val)
                copy_style(st, cell)
            r += 1
        last_data = r - 1
        if s["total"]:
            for c, (v, st) in enumerate(s["total"], start=1):
                col = get_column_letter(c)
                val = f"=SUM({col}{first}:{col}{last_data})" if isinstance(v, str) and v.startswith("=SUM(") else v
                cell = ws.cell(row=r, column=c, value=val)
                copy_style(st, cell)
            r += 1
        if s["check"]:
            for c, (v, st) in enumerate(s["check"], start=1):
                val = widen(v, old_last, new_last) if isinstance(v, str) and v.startswith("=") else v
                cell = ws.cell(row=r, column=c, value=val)
                copy_style(st, cell)
            r += 1
        r += 1


def main(csv_path: Path = CSV, xlsx_path: Path = XLSX) -> dict:
    df = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
    wb = openpyxl.load_workbook(xlsx_path)
    data = wb["Data"]
    old_last = data.max_row
    cols = [c.value for c in data[1]]
    assert cols == list(df.columns), f"Data sheet columns differ from the CSV: {cols} vs {list(df.columns)}"
    new_last = rewrite_rows(data, df.values.tolist())
    readable = df.sort_values(["date", "id"], kind="stable")
    rows, links = [], []
    for r in readable.itertuples(index=False):
        d = r._asdict()
        rows.append([d["id"], d["date"], COUNTRY.get(d["country"], d["country"]), "Yes" if d["country"] in IN_SCOPE else "No", d["speaker_name"], d["speaker_role"],
                     d["speaker_bloc"], label(d["channel"]), d["event_context"], d["summary_en"], label(d["minerals"]), label(d["themes"]), label(d["stance_china"]),
                     label(d["stance_us"]), d["quote_original"], d["quote_en"], d["language"], d["source_name"], label(d["source_type"]), label(d["verification"]), d["notes"]])
        links.append(d["source_url"] or None)
    st = wb["Statements"]
    assert [c.value for c in st[1]] == READABLE, "Statements sheet header changed"
    rewrite_rows(st, rows, hyperlinks={READABLE.index("Source") + 1: links})
    rebuild_coverage(wb["Coverage"], df, old_last, new_last)
    rd = wb["README"]
    for row in rd.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = widen(c.value, old_last, new_last)
    note = f"Regenerated from political_statements_minerals.csv by pipeline/scripts/statements_workbook.py on {date.today().isoformat()}: {len(df)} records."
    found = False
    for row in rd.iter_rows():
        if isinstance(row[0].value, str) and row[0].value.startswith("Regenerated from"):
            row[0].value, found = note, True
    if not found:
        rd.cell(row=rd.max_row + 2, column=1, value=note)
    wb.save(xlsx_path)
    return {"records": int(len(df)), "data_rows": new_last - 1, "workbook": str(xlsx_path.relative_to(ROOT)), "countries": int(df["country"].nunique())}


if __name__ == "__main__":
    args = [Path(a) for a in sys.argv[1:]]
    print(main(*args))
